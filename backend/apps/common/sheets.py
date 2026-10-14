"""Excel format files and uploads for the "Edit many" screens.

The format (template) is an .xlsx with one heading row, two example rows and
drop-down lists, plus a "How to fill" sheet. An upload is read back into plain
rows keyed by field name; the screen shows them for checking before anything
is saved (saving goes through bulk.py like typed rows).
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from decimal import Decimal

from django.http import HttpResponse

from apps.common.exceptions import BusinessError

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass
class SheetColumn:
    key: str
    label: str
    note: str
    required: bool = False
    choices: list[str] = field(default_factory=list)
    width: int = 16
    aliases: list[str] = field(default_factory=list)


@dataclass
class SheetFormat:
    title: str  # sheet and file name, e.g. "Items"
    columns: list[SheetColumn]
    examples: list[list]
    match_note: str


YES_NO = ["Yes", "No"]


def _units():
    from apps.products.models import Unit

    return [label for _, label in Unit.choices]


def _terms():
    from apps.parties.models import CreditTerms

    return [label for _, label in CreditTerms.choices]


def products_format() -> SheetFormat:
    return SheetFormat(
        title="Items",
        columns=[
            SheetColumn("name", "Item name", "Name of the item.", required=True, width=30, aliases=["name", "item", "product"]),
            SheetColumn("sku_code", "Item code", "Short code, different for every item (e.g. RICE-25). Used to find items already in the app.",
                        required=True, width=16, aliases=["sku", "code", "item code", "sku code"]),
            SheetColumn("unit", "Unit", "Pick from the list.", choices=_units(), width=10),
            SheetColumn("purchase_price", "Buying price", "Price you pay, without VAT.", width=14, aliases=["cost", "cost price", "purchase price"]),
            SheetColumn("selling_price", "Selling price", "Price you sell at, without VAT.", width=14, aliases=["sale price", "price", "rate"]),
            SheetColumn("tax_rate", "VAT %", "13 for VAT items, 0 for non-VAT items.", width=8, aliases=["vat", "tax", "tax %", "tax rate"]),
            SheetColumn("reorder_level", "Low stock alert", "Warn when stock falls to this quantity.", width=16, aliases=["reorder level", "reorder"]),
            SheetColumn("is_active", "In use", "Yes or No. Leave empty for Yes.", choices=YES_NO, width=8, aliases=["active"]),
        ],
        examples=[
            ["Basmati Rice 25kg", "RICE-25", "Kg", 1800, 2100, 13, 10, "Yes"],
            ["Sunflower Oil 1L", "OIL-1L", "Ltr", 260, 330, 13, 40, "Yes"],
        ],
        match_note="If an item code is already in the app, that item is changed. Otherwise a new item is added.",
    )


def parties_format(kind: str) -> SheetFormat:
    customer = kind == "CUSTOMER"
    columns = [
        SheetColumn("name", "Name", "Full name of the business or person.", required=True, width=30, aliases=["customer", "vendor", "party", "party name"]),
        SheetColumn("phone", "Phone", "Mobile or landline.", width=14, aliases=["mobile", "phone no", "contact"]),
        SheetColumn("email", "Email", "Email address.", width=24),
        SheetColumn("pan_vat_no", "PAN/VAT No.", "PAN or VAT number.", width=14, aliases=["pan", "vat no", "pan no", "pan/vat"]),
        SheetColumn("address", "Address", "Address.", width=24),
        SheetColumn("credit_terms", "Pays within", "Pick from the list. Choose Custom to type your own days.", choices=_terms(), width=12,
                    aliases=["credit terms", "terms"]),
        SheetColumn("credit_days", "Days (if Custom)", "Only used when 'Pays within' is Custom.", width=12, aliases=["days", "credit days"]),
    ]
    if customer:
        columns.append(SheetColumn("credit_limit", "Max credit (0 = no limit)", "Largest amount the customer may owe. 0 = no limit.", width=18,
                                   aliases=["credit limit", "max credit", "limit"]))
    columns.append(SheetColumn("is_active", "In use", "Yes or No. Leave empty for Yes.", choices=YES_NO, width=8, aliases=["active"]))
    examples = [
        ["Sharma Kirana Pasal", "9800000001", "", "601234567", "Kathmandu", "30 Days", "", *([50000] if customer else []), "Yes"],
        ["Gurung Store", "9800000002", "gurung@example.com", "", "Pokhara", "Cash", "", *([0] if customer else []), "Yes"],
    ]
    return SheetFormat(
        title="Customers" if customer else "Vendors",
        columns=columns,
        examples=examples,
        match_note="If a name is already in the app, that one is changed. Otherwise a new one is added.",
    )


# ---------------------------------------------------------------------------
# Format file
# ---------------------------------------------------------------------------
def template_response(fmt: SheetFormat) -> HttpResponse:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = fmt.title
    head_fill = PatternFill("solid", fgColor="1E3A8A")
    need_fill = PatternFill("solid", fgColor="B91C1C")
    for i, col in enumerate(fmt.columns, start=1):
        cell = ws.cell(row=1, column=i, value=col.label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = need_fill if col.required else head_fill
        cell.alignment = Alignment(vertical="center")
        cell.comment = Comment(col.note + (" Must be filled in." if col.required else ""), "Accounting")
        ws.column_dimensions[get_column_letter(i)].width = col.width
        if col.choices:
            dv = DataValidation(type="list", formula1='"' + ",".join(col.choices) + '"', allow_blank=True)
            dv.error = "Pick one from the list."
            dv.errorTitle = col.label
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(i)}2:{get_column_letter(i)}{MAX_ROWS + 1}")
    for r, row in enumerate(fmt.examples, start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value if value != "" else None).font = Font(italic=True, color="64748B")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    help_ws = wb.create_sheet("How to fill")
    lines = [
        (f"How to fill the {fmt.title} sheet", True),
        ("", False),
        ("1. Fill one row for each " + fmt.title.lower().rstrip("s") + ". Start on row 2.", False),
        ("2. Delete the two grey example rows (or type over them).", False),
        ("3. Red headings must be filled in. Blue ones can be left empty.", False),
        ("4. Do not change the headings in row 1.", False),
        ("5. Save the file, then use 'Upload Excel' on the 'Edit many' screen.", False),
        ("6. Check the rows on screen, then press Save. Nothing is saved before that.", False),
        ("", False),
        (fmt.match_note, False),
        ("", False),
        ("Columns", True),
    ]
    for text, bold in lines:
        help_ws.append([text])
        if bold:
            help_ws.cell(row=help_ws.max_row, column=1).font = Font(bold=True, size=12 if help_ws.max_row == 1 else 11)
    for col in fmt.columns:
        choices = f" Choices: {', '.join(col.choices)}." if col.choices else ""
        help_ws.append([col.label + (" (must fill)" if col.required else ""), col.note + choices])
        help_ws.cell(row=help_ws.max_row, column=1).font = Font(bold=True)
    help_ws.column_dimensions["A"].width = 28
    help_ws.column_dimensions["B"].width = 90

    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type=XLSX_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{fmt.title.lower()}-format.xlsx"'
    return response


# ---------------------------------------------------------------------------
# Upload
# ---------------------------------------------------------------------------
def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        value = Decimal(str(value))
    if isinstance(value, Decimal):
        value = value.normalize()
        return format(value, "f")
    if hasattr(value, "isoformat"):
        return value.isoformat()[:10]
    return str(value).strip()


def _read_rows(upload) -> list[list]:
    name = (upload.name or "").lower()
    if upload.size > MAX_BYTES:
        raise BusinessError("The file is too big. Keep it under 5 MB.", code="FILE_TOO_BIG")
    data = upload.read()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        return list(csv.reader(io.StringIO(text)))
    if name.endswith(".xls"):
        raise BusinessError("This is an old Excel file (.xls). In Excel use File > Save As > Excel Workbook (.xlsx), then upload that.",
                            code="OLD_EXCEL")
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as err:  # noqa: BLE001 - any unreadable file
        raise BusinessError("Could not open this file. Upload an Excel (.xlsx) or CSV file.", code="BAD_FILE") from err
    ws = wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def parse_upload(fmt: SheetFormat, upload) -> dict:
    if upload is None:
        raise BusinessError("Choose a file to upload.", code="NO_FILE")
    raw = _read_rows(upload)

    def norm(v) -> str:
        return " ".join(_text(v).lower().replace("*", "").split())

    lookup = {}
    for col in fmt.columns:
        for name in [col.label, col.key, *col.aliases]:
            lookup[norm(name)] = col.key

    # The heading row is the first row (of the first 10) that names a known column.
    header_at, mapping = None, {}
    for i, row in enumerate(raw[:10]):
        found = {j: lookup[norm(v)] for j, v in enumerate(row) if norm(v) in lookup}
        if found:
            header_at, mapping = i, found
            break
    if header_at is None:
        raise BusinessError(
            f"Could not find the headings. Download the format and keep row 1 as it is (e.g. '{fmt.columns[0].label}').",
            code="NO_HEADINGS",
        )
    missing = [c.label for c in fmt.columns if c.required and c.key not in mapping.values()]
    if missing:
        raise BusinessError(f"The file has no '{missing[0]}' column. Download the format and use its headings.", code="MISSING_COLUMN")

    examples = {tuple(norm(v) for v in ex[:2]) for ex in fmt.examples}
    rows, skipped_examples = [], 0
    for row in raw[header_at + 1:]:
        values = {key: _text(row[j]) for j, key in mapping.items() if j < len(row)}
        if not any(values.values()):
            continue
        first_two = tuple(norm(values.get(c.key, "")) for c in fmt.columns[:2])
        if first_two in examples:
            skipped_examples += 1
            continue
        rows.append(values)
    if not rows:
        raise BusinessError("The file has no rows to load.", code="EMPTY_FILE")
    if len(rows) > MAX_ROWS:
        raise BusinessError(f"The file has {len(rows)} rows. Upload at most {MAX_ROWS} at a time.", code="TOO_MANY_ROWS")
    labels = {c.key: c.label for c in fmt.columns}
    return {
        "rows": rows,
        "columns": [labels[k] for k in dict.fromkeys(mapping.values())],
        "skipped_examples": skipped_examples,
    }
