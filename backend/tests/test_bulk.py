from decimal import Decimal as D

from apps.parties.models import Party, PartyType
from apps.products.models import Product

from .conftest import client_for, make_user


def test_bulk_products_create_and_update(api, product_factory):
    p = product_factory(name="Rice", cost=D("100"), price=D("120"))
    rows = [
        {"id": p.pk, "selling_price": "130.00", "tax_rate": "0"},
        {"name": "Dal", "sku_code": "dal-1", "unit": "KG", "purchase_price": "140", "selling_price": "165"},
        {"name": "Oil", "sku_code": "OIL-1", "unit": "LTR", "is_active": False},
    ]
    r = api.post("/api/products/bulk/", {"rows": rows}, format="json")
    assert r.status_code == 200, r.content
    assert r.json()["data"] == {"created": 2, "updated": 1}
    assert r.json()["message"] == "Saved: 2 new items added, 1 item changed."
    p.refresh_from_db()
    assert p.selling_price == D("130.00") and p.tax_rate == D("0")
    assert Product.objects.get(sku_code="DAL-1").selling_price == D("165")
    assert Product.objects.get(sku_code="OIL-1").is_active is False
    # Unchanged rows are not counted.
    assert api.post("/api/products/bulk/", {"rows": [{"id": p.pk, "selling_price": "130"}]}, format="json").json()["data"]["updated"] == 0


def test_bulk_products_all_or_nothing(api, product_factory):
    p = product_factory(name="Rice")
    rows = [
        {"id": p.pk, "selling_price": "999"},
        {"name": "A", "sku_code": "DUP"},
        {"name": "B", "sku_code": "dup"},
        {"name": "", "sku_code": "X1", "selling_price": "-5"},
    ]
    r = api.post("/api/products/bulk/", {"rows": rows}, format="json")
    assert r.status_code == 400
    errors = r.json()["error"]["details"]["rows"]
    assert set(errors) == {"2", "3"}
    assert "sku_code" in errors["2"] and {"name", "selling_price"} <= set(errors["3"])
    p.refresh_from_db()
    assert p.selling_price != D("999")
    assert not Product.objects.filter(sku_code="DUP").exists()


def test_bulk_customers(api, customer):
    rows = [
        {"id": customer.pk, "phone": "9800000000", "credit_terms": "DAYS_45", "credit_limit": "50000"},
        {"name": "New Shop", "phone": "9811111111", "credit_terms": "CASH"},
    ]
    r = api.post("/api/customers/bulk/", {"rows": rows}, format="json")
    assert r.status_code == 200, r.content
    customer.refresh_from_db()
    assert customer.credit_days == 45 and customer.credit_limit == D("50000") and customer.phone == "9800000000"
    assert Party.objects.get(name="New Shop").type == PartyType.CUSTOMER


def test_bulk_vendors_cannot_touch_customers(api, customer):
    r = api.post("/api/vendors/bulk/", {"rows": [{"id": customer.pk, "name": "Hijack"}]}, format="json")
    assert r.status_code == 400
    customer.refresh_from_db()
    assert customer.name != "Hijack"


def test_bulk_needs_manage_permission(db, customer):
    from apps.users.models import Role

    staff = client_for(make_user(Role.STAFF, "bulk_staff"))
    assert staff.post("/api/customers/bulk/", {"rows": [{"name": "X"}]}, format="json").status_code == 403
    assert staff.post("/api/products/bulk/", {"rows": [{"name": "X", "sku_code": "X"}]}, format="json").status_code == 403


# ---------------------------------------------------------------------------
# Excel format and upload
# ---------------------------------------------------------------------------
def _xlsx(rows):
    import io

    from django.core.files.uploadedfile import SimpleUploadedFile
    from openpyxl import Workbook

    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile("items.xlsx", buf.getvalue())


def test_format_file_downloads(api):
    import io

    from openpyxl import load_workbook

    for url, first in [("/api/products/bulk/template/", "Item name"), ("/api/customers/bulk/template/", "Name"), ("/api/vendors/bulk/template/", "Name")]:
        r = api.get(url)
        assert r.status_code == 200 and r["Content-Type"].startswith("application/vnd.openxmlformats")
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.worksheets[0].cell(1, 1).value == first
        assert wb.sheetnames[1] == "How to fill"
    headers = [c.value for c in load_workbook(io.BytesIO(api.get("/api/vendors/bulk/template/").content)).worksheets[0][1]]
    assert "Max credit (0 = no limit)" not in headers


def test_upload_reads_rows_and_skips_examples(api):
    f = _xlsx([
        ["Item name", "Item code", "Unit", "Buying price", "Selling price", "VAT %"],
        ["Basmati Rice 25kg", "RICE-25", "Kg", 1800, 2100, 13],  # the example row
        ["Sugar 1kg", "SUG-1", "kg", 95.5, 110, 0],
        [None, None, None, None, None, None],
        ["Salt", "SALT", "Pcs", "", 25, 13],
    ])
    r = api.post("/api/products/bulk/upload/", {"file": f}, format="multipart")
    assert r.status_code == 200, r.content
    d = r.json()["data"]
    assert d["skipped_examples"] == 1
    assert d["rows"][0] == {"name": "Sugar 1kg", "sku_code": "SUG-1", "unit": "kg", "purchase_price": "95.5", "selling_price": "110", "tax_rate": "0"}
    assert len(d["rows"]) == 2
    # Nothing was saved.
    from apps.products.models import Product

    assert not Product.objects.filter(sku_code="SUG-1").exists()


def test_upload_accepts_csv_and_aliases(api):
    from django.core.files.uploadedfile import SimpleUploadedFile

    csv = "Customer,Mobile,Credit Terms,Credit Limit\nRam Traders,9800000000,30 Days,5000\n".encode()
    r = api.post("/api/customers/bulk/upload/", {"file": SimpleUploadedFile("c.csv", csv)}, format="multipart")
    assert r.status_code == 200, r.content
    assert r.json()["data"]["rows"] == [{"name": "Ram Traders", "phone": "9800000000", "credit_terms": "30 Days", "credit_limit": "5000"}]


def test_upload_explains_bad_files(api):
    from django.core.files.uploadedfile import SimpleUploadedFile

    no_heads = api.post("/api/products/bulk/upload/", {"file": _xlsx([["a", "b"], ["c", "d"]])}, format="multipart")
    assert no_heads.status_code == 400 and no_heads.json()["error"]["code"] == "NO_HEADINGS"
    missing = api.post("/api/products/bulk/upload/", {"file": _xlsx([["Item name"], ["Rice"]])}, format="multipart")
    assert missing.json()["error"]["code"] == "MISSING_COLUMN"
    old = api.post("/api/products/bulk/upload/", {"file": SimpleUploadedFile("x.xls", b"junk")}, format="multipart")
    assert old.json()["error"]["code"] == "OLD_EXCEL"
    junk = api.post("/api/products/bulk/upload/", {"file": SimpleUploadedFile("x.xlsx", b"junk")}, format="multipart")
    assert junk.json()["error"]["code"] == "BAD_FILE"
